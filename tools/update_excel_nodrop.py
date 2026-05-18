"""X7 강화 시스템 Excel 업데이트 — No-drop 시스템 반영
- 확률표: 새 확률값, 실패패널티 "없음"으로 변경
- 기대값: 기하분포 E(k)=1/p 공식으로 재계산
- 델피나드: 천장 통계 새 확률로 재계산
- 노강소모: MC 재시뮬레이션 (no-drop)
- 요약: 모든 섹션 동기화
"""
import math, random, shutil, time
import openpyxl
from openpyxl import load_workbook

SRC  = r'C:\Users\hoy5343\Downloads\X7_강화시스템_티어별_수정.xlsx'
DST  = r'C:\Users\hoy5343\Downloads\X7_강화시스템_티어별_v2_노드롭.xlsx'

shutil.copy2(SRC, DST)
print(f"복사 완료: {DST}")

# ──────────────────────────────────────────────────────────────
# 상수
# ──────────────────────────────────────────────────────────────
NEW_PROBS = {
    1:  [1.00, 1.00, 1.00, 1.00, 1.00, 1.00, 1.00],
    2:  [1.00, 1.00, 1.00, 1.00, 1.00, 1.00, 1.00],
    3:  [1.00, 1.00, 1.00, 1.00, 1.00, 1.00, 1.00],
    4:  [0.80, 0.75, 0.70, 0.65, 0.60, 0.55, 0.50],
    5:  [0.65, 0.60, 0.55, 0.50, 0.45, 0.40, 0.35],
    6:  [0.50, 0.45, 0.40, 0.35, 0.30, 0.25, 0.20],
    7:  [0.40, 0.35, 0.30, 0.25, 0.20, 0.15, 0.10],
    8:  [0.25, 0.22, 0.19, 0.16, 0.13, 0.11, 0.09],
    9:  [0.16, 0.14, 0.12, 0.10, 0.08, 0.06, 0.05],
    10: [0.10, 0.09, 0.08, 0.06, 0.05, 0.04, 0.03],
}

ENERGY = {
    4:  {'epf': 20.0, 'boost': 2.0},
    5:  {'epf': 15.0, 'boost': 1.5},
    6:  {'epf': 10.0, 'boost': 1.2},
    7:  {'epf':  8.0, 'boost': 1.0},
    8:  {'epf':  6.0, 'boost': 1.0},
    9:  {'epf':  5.0, 'boost': 1.0},
    10: {'epf':  4.0, 'boost': 1.0},
}

TIER_LABELS = ['T1', 'T2', 'T3', 'T4', 'T5', 'T6', 'T7']

# ──────────────────────────────────────────────────────────────
# 계산 함수
# ──────────────────────────────────────────────────────────────
def compute_step_ev():
    """단계별 기대 시도수 E(k) = 1/p(k) (no-drop 기하분포)"""
    ev = {}
    for lvl in range(1, 11):
        ev[lvl] = [1.0 / p for p in NEW_PROBS[lvl]]
    return ev

def compute_cumulative_ev(step_ev):
    """누적 기대 시도수 (1~k 합산)"""
    cum = {}
    running = [0.0] * 7
    for lvl in range(1, 11):
        for ti in range(7):
            running[ti] += step_ev[lvl][ti]
        cum[lvl] = running[:]
    return cum

def compute_ceiling_p(lvl, ti):
    """천장 도달 확률 P(천장)"""
    pb  = NEW_PROBS[lvl][ti]
    epf = ENERGY[lvl]['epf']
    bst = ENERGY[lvl]['boost']
    N   = math.ceil(100.0 / epf)
    p_not = 1.0
    for k in range(N):
        pk = min(1.0, pb + k * bst / 100.0)
        p_not *= (1.0 - pk)
    return p_not

def mc_nodrop(tier_idx, n_sim=2000):
    """No-drop MC: pre[k] = +k 아이템 1개 만들기 위한 노강 기댓값"""
    rng = random.Random(42 + tier_idx * 137 + n_sim)
    pre = {0: 1.0}
    for target in range(1, 11):
        total = 0.0
        for _ in range(n_sim):
            level, used, attempts = 0, 1.0, 3
            pity  = {k: [0.0, 0.0] for k in range(4, 11)}
            guard = 0
            while level < target and guard < 300_000:
                guard += 1
                tlvl = level + 1
                pb   = NEW_PROBS[tlvl][tier_idx]
                if tlvl >= 4 and tlvl in ENERGY:
                    ba, ea = pity[tlvl]
                    p_eff = 1.0 if ea >= 100.0 else min(1.0, pb + ba / 100.0)
                else:
                    p_eff = pb
                attempts -= 1
                if rng.random() < p_eff:
                    level += 1
                    if tlvl >= 4:
                        pity[tlvl] = [0.0, 0.0]
                else:
                    if tlvl >= 4 and tlvl in ENERGY:
                        cfg = ENERGY[tlvl]
                        pity[tlvl][0] += cfg['boost']
                        pity[tlvl][1]  = min(100.0, pity[tlvl][1] + cfg['epf'])
                    # 단계 하락 없음
                if attempts == 0 and level < target:
                    used    += pre.get(level, 1.0)
                    attempts = 3
            total += used
        pre[target] = total / n_sim
    return pre

def nogang_label(val):
    if val <= 2:   return '✅ 소량'
    if val <= 10:  return '🟡 약 10개↓'
    if val <= 50:  return '🟠 ~50개'
    if val <= 200: return '🔴 ~100개'
    if val <= 2000:return '⛔ 수백~수천'
    return '💀 수만개+'

# ──────────────────────────────────────────────────────────────
# 사전 계산
# ──────────────────────────────────────────────────────────────
print("\n[1/4] 기대값 계산...")
step_ev = compute_step_ev()
cum_ev  = compute_cumulative_ev(step_ev)

print("[2/4] 천장 통계 계산...")
ceil_p   = {lvl: [compute_ceiling_p(lvl, ti) for ti in range(7)] for lvl in range(4, 11)}
succ_p   = {lvl: [1.0 - ceil_p[lvl][ti]      for ti in range(7)] for lvl in range(4, 11)}

print("[3/4] MC 노강소모 시뮬레이션 (n=2000 × 7티어)...")
t0 = time.time()
mc_all = {}
for ti in range(7):
    mc_all[ti] = mc_nodrop(ti, n_sim=2000)
    print(f"  {TIER_LABELS[ti]} 완료  +10={mc_all[ti][10]:.0f}개")
print(f"  MC 완료 ({time.time()-t0:.1f}s)")

# ──────────────────────────────────────────────────────────────
# 엑셀 쓰기
# ──────────────────────────────────────────────────────────────
print("\n[4/4] Excel 업데이트...")
wb = load_workbook(DST)

# ──────────────────────────────────────────────
# 확률표 시트
# ──────────────────────────────────────────────
ws = wb['확률표']
# 설명 행 업데이트
ws.cell(3, 2).value = (
    '🔵 파란 테두리 셀 = 직접 수정 가능 | 실패 시 단계 유지 (하락 없음) | T4 = 기준값'
)
# 확률값 + 실패패널티 업데이트 (row 8~17 = level 1~10)
penalties = {
    1: '없음', 2: '없음', 3: '없음',
    4: '없음', 5: '없음', 6: '없음',
    7: '없음', 8: '없음', 9: '없음', 10: '없음',
}
for idx, lvl in enumerate(range(1, 11)):
    row = 8 + idx
    for ti in range(7):
        ws.cell(row, 3 + ti).value = NEW_PROBS[lvl][ti]
    ws.cell(row, 10).value = penalties[lvl]

print("  확률표 완료")

# ──────────────────────────────────────────────
# 기대값 시트
# ──────────────────────────────────────────────
ws2 = wb['기대값']
# 설명 업데이트
ws2.cell(3, 2).value = (
    '기하분포: E(k) = 1/p(k)  |  실패 시 단계 유지 (하락 없음)  |  T4 = 기준값'
)
ws2.cell(8, 2).value = '■  단계별 기대 시도수 (해당 레벨→다음 레벨 평균 시도)'

# 단계별 기대 시도수 (row 9~18 = level 1~10)
for idx, lvl in enumerate(range(1, 11)):
    row = 9 + idx
    for ti in range(7):
        ws2.cell(row, 3 + ti).value = round(step_ev[lvl][ti], 6)
    # T7/T1 배율 (col 10)
    ratio = step_ev[lvl][6] / step_ev[lvl][0] if step_ev[lvl][0] > 0 else 1.0
    ws2.cell(row, 10).value = round(ratio, 6)

# 누적 기대 시도수 (row 22~31 = level 1~10)
for idx, lvl in enumerate(range(1, 11)):
    row = 22 + idx
    for ti in range(7):
        ws2.cell(row, 3 + ti).value = round(cum_ev[lvl][ti], 6)
    ratio = cum_ev[lvl][6] / cum_ev[lvl][0] if cum_ev[lvl][0] > 0 else 1.0
    ws2.cell(row, 10).value = round(ratio, 6)

print("  기대값 완료")

# ──────────────────────────────────────────────
# 델피나드 시트
# ──────────────────────────────────────────────
ws3 = wb['델피나드']
# 천장 도달 확률 (row 26~32 = level 4~10)
for idx, lvl in enumerate(range(4, 11)):
    row_c = 26 + idx
    row_s = 35 + idx
    for ti in range(7):
        ws3.cell(row_c, 3 + ti).value = round(ceil_p[lvl][ti], 10)
        ws3.cell(row_s, 3 + ti).value = round(succ_p[lvl][ti], 10)

print("  델피나드 완료")

# ──────────────────────────────────────────────
# 노강소모 시트
# ──────────────────────────────────────────────
ws4 = wb['노강소모']
# 헤더 업데이트
ws4.cell(2, 2).value = '🔧  노강 아이템 기대 소모량  (Monte Carlo n=2,000 · 천장 시스템 포함 · 단계 하락 없음)'
ws4.cell(3, 2).value = '조건: 강화 3회 소진 → 현재 강화 단계 동일 아이템 1개 소모 복구  |  실패 시 단계 유지 (하락 없음)'

# 데이터 (row 7~16 = level 1~10)
for idx, lvl in enumerate(range(1, 11)):
    row = 7 + idx
    for ti in range(7):
        val = round(mc_all[ti][lvl])
        ws4.cell(row, 3 + ti).value = val
    # 규모 판정 (T4 기준)
    ws4.cell(row, 10).value = nogang_label(round(mc_all[3][lvl]))

# 요약 텍스트 업데이트 (row 23~24)
t4_10  = round(mc_all[3][10])
t7_10  = round(mc_all[6][10])
ws4.cell(23, 2).value = 'T4 +10'
ws4.cell(23, 3).value = f'기대 소모 ≈ {t4_10:,}개  (천장 시스템으로 소모량 현실화)'
ws4.cell(24, 2).value = 'T7 +10'
ws4.cell(24, 3).value = f'기대 소모 ≈ {t7_10:,}개  (극한 희소성 — 서버 이벤트·보호권 필수 고려)'

print("  노강소모 완료")

# ──────────────────────────────────────────────
# 요약 시트
# ──────────────────────────────────────────────
ws5 = wb['요약']
# 확률 (row 7~13 = level 4~10)
difficulty = {4:'보통', 5:'높음', 6:'높음', 7:'매우 높음', 8:'매우 높음', 9:'극한', 10:'극한+'}
for idx, lvl in enumerate(range(4, 11)):
    row = 7 + idx
    for ti in range(7):
        ws5.cell(row, 3 + ti).value = NEW_PROBS[lvl][ti]
    ws5.cell(row, 10).value = difficulty[lvl]

# 누적 기대 시도수 (row 18~24 = level 4~10)
for idx, lvl in enumerate(range(4, 11)):
    row = 18 + idx
    for ti in range(7):
        ws5.cell(row, 3 + ti).value = round(cum_ev[lvl][ti], 6)
    ratio = cum_ev[lvl][6] / cum_ev[lvl][0] if cum_ev[lvl][0] > 0 else 1.0
    ws5.cell(row, 10).value = round(ratio, 4)

# 노강소모 (row 29~35 = level 4~10)
for idx, lvl in enumerate(range(4, 11)):
    row = 29 + idx
    for ti in range(7):
        ws5.cell(row, 3 + ti).value = round(mc_all[ti][lvl])
    ws5.cell(row, 10).value = nogang_label(round(mc_all[3][lvl]))

print("  요약 완료")

# ──────────────────────────────────────────────
wb.save(DST)
print(f"\n✅ 저장 완료: {DST}")

# ──────────────────────────────────────────────
# 결과 요약 출력
# ──────────────────────────────────────────────
print("\n====== 주요 수치 요약 ======")
print(f"\n[확률표] T4 기준")
for lvl in [4,5,6,7,8,9,10]:
    print(f"  +{lvl}: {NEW_PROBS[lvl][3]*100:.0f}%")

print(f"\n[기대값 누적] T4 (+k까지 총 시도)")
for lvl in [4,5,6,7,8,9,10]:
    print(f"  +{lvl}: {cum_ev[lvl][3]:.1f}회")

print(f"\n[노강소모] T4 / T7")
for lvl in [4,5,6,7,8,9,10]:
    print(f"  +{lvl}: T4={round(mc_all[3][lvl]):>6,}개  T7={round(mc_all[6][lvl]):>8,}개")
